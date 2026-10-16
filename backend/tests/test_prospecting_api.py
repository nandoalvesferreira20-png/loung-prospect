from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from threading import Event
from time import monotonic, sleep
import hashlib
import json
import sqlite3
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.config import Settings
from backend.main import create_app
from backend.prospecting.models import SearchRequest
from backend.prospecting.storage import SelectionStore, StorageError
from backend.tests.test_api import database  # noqa: F401 (shared temporary fixture)
from core.database.lead_repository import LeadRepository
from core.database.schema import initialize_database
from core.providers.google_places import GooglePlacesProvider
from core.providers.models import LeadCandidate
from core.prospecting.service import prospect

QUERY={"niche":"Clínicas odontológicas","location":"São Paulo, SP","quantity":3}


def place(index, **extra):
    return {"id":f"synthetic-{index}","displayName":{"text":f"Teste busca {index}"},"formattedAddress":f"Rua de teste {index}","nationalPhoneNumber":"(11) 99999-0000",**extra}


def transport(payload, headers, timeout):
    data={"places":[place(3)]} if "pageToken" in payload else {"places":[place(1),place(2)],"nextPageToken":"unit-page-2"}
    return 200,json.dumps(data).encode()


@contextmanager
def api(path,monkeypatch,*,enabled=True,send=transport,options=None):
    # Neither .env nor a live network transport is used by any test.
    monkeypatch.setenv("GOOGLE_PLACES_API_KEY","unit-test-not-a-real-key")
    monkeypatch.setattr("core.providers.environment.load_environment",lambda:pytest.fail("Must not load .env in tests"))
    monkeypatch.setattr("backend.prospecting.jobs.load_environment",lambda:pytest.fail("Must not load .env in tests"))
    monkeypatch.setattr("core.providers.google_places._post",lambda *args:pytest.fail("Live Places transport forbidden"))
    with TestClient(create_app(Settings(path,enabled),provider_factory=lambda:GooglePlacesProvider(transport=send,timeout=.1,max_requests=3),job_options=options)) as client:
        yield client


def wait(client,job_id):
    end=monotonic()+5
    while monotonic()<end:
        response=client.get(f"/v1/prospecting/search/{job_id}")
        job=response.json()
        if job["status"] in {"completed","failed","cancelled"}:return job
        sleep(.01)
    pytest.fail("Job did not finish")


def start(client,query=None):
    response=client.post("/v1/prospecting/search",json=query or QUERY)
    assert response.status_code==202,response.text
    return response.json()["id"]


def test_search_paginate_export_and_no_automatic_write(database,monkeypatch):
    before=database.read_bytes()
    with api(database,monkeypatch) as client:
        job_id=start(client);job=wait(client,job_id)
        assert job["status"]=="completed" and job["found"]==3
        assert job["requestsMade"]==2 and job["pagesFetched"]==2
        page=client.get(f"/v1/prospecting/search/{job_id}/results?pageSize=2").json()
        assert page["total"]==3 and len(page["items"])==2
        assert client.get(f"/v1/prospecting/search/{job_id}/results?pageSize=2&page=2").json()["items"][0]["id"]==3
        assert client.get(f"/v1/prospecting/search/{job_id}/results?page=0").status_code==422
        csv=client.get(f"/v1/prospecting/search/{job_id}/export")
        assert csv.status_code==200 and "Teste busca 1" in csv.content.decode("utf-8-sig")
        excel=client.get(f"/v1/prospecting/search/{job_id}/export?format=xlsx")
        assert load_workbook(BytesIO(excel.content)).active.max_row==4
    assert database.read_bytes()==before


def test_import_transaction_backup_dedupe_and_existing_unchanged(database,monkeypatch):
    with sqlite3.connect(database) as connection:
        before=connection.execute("SELECT * FROM leads ORDER BY id").fetchall()
    with api(database,monkeypatch) as client:
        job_id=start(client);wait(client,job_id)
        request={"searchId":job_id,"resultIds":[1,2,2]}
        response=client.post("/v1/leads/import",json=request)
        assert response.json()=={"inserted":2,"duplicates":0},response.text
        assert client.post("/v1/leads/import",json=request).json()=={"inserted":0,"duplicates":2}
        results=client.get(f"/v1/prospecting/search/{job_id}/results").json()["items"]
        assert [r["alreadySaved"] for r in results]==[True,True,False]
        assert client.get("/v1/leads").json()["total"]==7
    with sqlite3.connect(database) as connection:
        assert connection.execute("SELECT * FROM leads WHERE id<=5 ORDER BY id").fetchall()==before
    backups=list((database.parent/"backups").glob("*.db"))
    assert len(backups)==1
    with sqlite3.connect(backups[0]) as backup:
        assert backup.execute("SELECT COUNT(*) FROM leads").fetchone()[0]==5


def test_write_first_in_snapshot_copy_never_source(database,tmp_path,monkeypatch):
    source_hash=hashlib.sha256(database.read_bytes()).hexdigest()
    copy=tmp_path/"desktop-copy.db"
    with sqlite3.connect(database) as source,sqlite3.connect(copy) as destination:source.backup(destination)
    with api(copy,monkeypatch) as client:
        job_id=start(client);wait(client,job_id)
        assert client.post("/v1/leads/import",json={"searchId":job_id,"resultIds":[1]}).json()["inserted"]==1
    assert hashlib.sha256(database.read_bytes()).hexdigest()==source_hash


@pytest.mark.parametrize("change",[{"quantity":101},{"quantity":True},{"niche":" "},{"location":""},{"minScore":101},{"extra":"x"}])
def test_validation(database,monkeypatch,change):
    with api(database,monkeypatch) as client:
        assert client.post("/v1/prospecting/search",json={**QUERY,**change}).status_code==422


def test_failures_preserve_partial_results_and_hide_secrets(database,monkeypatch):
    def failure(payload,headers,timeout):
        return (500,b'{"secret":"unit-sensitive-value"}') if "pageToken" in payload else transport(payload,headers,timeout)
    with api(database,monkeypatch,send=failure) as client:
        job_id=start(client);job=wait(client,job_id)
        assert job["status"]=="failed" and job["found"]==2
        assert "unit-sensitive-value" not in json.dumps(job)
        assert len(client.get(f"/v1/prospecting/search/{job_id}/results").json()["items"])==2


def test_cancel_concurrency_timeout_and_no_write(database,monkeypatch):
    entered=Event();release=Event()
    def slow(payload,headers,timeout):
        entered.set();release.wait(2);return transport(payload,headers,timeout)
    before=database.read_bytes()
    with api(database,monkeypatch,send=slow) as client:
        job_id=start(client);assert entered.wait(1)
        assert client.post("/v1/prospecting/search",json=QUERY).status_code==429
        assert client.post("/v1/leads/import",json={"searchId":job_id,"resultIds":[1]}).status_code==409
        assert client.post(f"/v1/prospecting/search/{job_id}/cancel").json()["cancelRequested"]
        release.set();assert wait(client,job_id)["status"]=="cancelled"
    assert database.read_bytes()==before
    with api(database,monkeypatch,options={"timeout":0}) as client:
        assert wait(client,start(client))["status"]=="failed"


def test_disabled_writes_unknown_ids_and_expiration(database,monkeypatch):
    before=database.read_bytes()
    with api(database,monkeypatch,enabled=False) as client:
        job_id=start(client);assert not wait(client,job_id)["writesEnabled"]
        assert client.post("/v1/leads/import",json={"searchId":job_id,"resultIds":[1]}).status_code==503
        assert client.post("/v1/leads/import",json={"searchId":job_id,"resultIds":[99]}).status_code==422
        assert client.get("/v1/prospecting/search/absent").status_code==404
        client.app.state.prospecting.retention=0
        assert client.get(f"/v1/prospecting/search/{job_id}").status_code==404
    assert database.read_bytes()==before


def test_spreadsheet_formula_safety(database,monkeypatch):
    def formula(payload,headers,timeout):return 200,json.dumps({"places":[place(1,displayName={"text":"=HYPERLINK(unsafe)"})]}).encode()
    with api(database,monkeypatch,send=formula) as client:
        job_id=start(client);wait(client,job_id)
        assert "'=HYPERLINK" in client.get(f"/v1/prospecting/search/{job_id}/export").content.decode("utf-8-sig")


def test_backup_failure_rolls_back_everything(database,monkeypatch):
    store=SelectionStore(database,True)
    before=database.read_bytes()
    def fail():raise OSError("unit-only")
    monkeypatch.setattr(store,"_backup",fail)
    with pytest.raises(StorageError):store.save([{"empresa":"Teste","provider":"google_places","provider_place_id":"not-real"}])
    assert database.read_bytes()==before


def test_late_invalid_row_rolls_back_earlier_inserts(database):
    before=database.read_bytes()
    store=SelectionStore(database,True)
    with pytest.raises(StorageError):
        store.save([{"empresa":"Teste transação","provider":"unit","provider_place_id":"rollback"},{"empresa":" "}])
    assert database.read_bytes()==before


def test_cancel_preserves_completed_page_and_allows_selected_import(database,monkeypatch):
    entered,release=Event(),Event()
    def partial(payload,headers,timeout):
        if "pageToken" not in payload:
            return 200,json.dumps({"places":[place(1)],"nextPageToken":"next"}).encode()
        entered.set();release.wait(2)
        return 200,json.dumps({"places":[place(2)]}).encode()
    with api(database,monkeypatch,send=partial) as client:
        job_id=start(client);assert entered.wait(1)
        assert client.get(f"/v1/prospecting/search/{job_id}").json()["found"]==1
        client.post(f"/v1/prospecting/search/{job_id}/cancel")
        release.set()
        assert wait(client,job_id)["status"]=="cancelled"
        assert client.get(f"/v1/prospecting/search/{job_id}/results").json()["total"]==1
        assert client.post("/v1/leads/import",json={"searchId":job_id,"resultIds":[1]}).json()["inserted"]==1


def test_shared_streaming_engine_reports_actual_requests(monkeypatch):
    monkeypatch.setenv("GOOGLE_PLACES_API_KEY","synthetic-only")
    class Sink:
        def create_lead_if_new(self, data): return 1
    summary=prospect(city="São Paulo",segment="Teste",limit=3,
        provider=GooglePlacesProvider(transport=transport),repository=Sink(),stream_results=True)
    assert summary.received==3 and summary.requests_made==2 and summary.pages_fetched==2


def test_unavailable_workspace_never_calls_provider(tmp_path):
    def forbidden():pytest.fail("Provider must not start for unavailable database")
    with TestClient(create_app(Settings(tmp_path/"absent.db"),provider_factory=forbidden)) as client:
        assert client.post("/v1/prospecting/search",json=QUERY).status_code==503
    assert not (tmp_path/"absent.db").exists()
