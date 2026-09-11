from unittest.mock import Mock

import pandas as pd
import pytest
from openpyxl import load_workbook

from core.exporter import export_excel


def test_excel_roundtrip_preserves_columns_values_and_no_index(tmp_path, lead_factory, columns):
    row = lead_factory()
    output = tmp_path / "nested" / "leads.xlsx"
    result = export_excel(pd.DataFrame([row]), str(output), log=Mock())
    assert result == output
    restored = pd.read_excel(output, keep_default_na=False)
    assert list(restored.columns) == columns
    assert restored.to_dict("records") == [row]
    workbook = load_workbook(output)
    try:
        assert workbook.active.max_column == len(columns)
        assert workbook.active.max_row == 2
    finally:
        workbook.close()


def test_export_error_is_logged_and_reraised(tmp_path):
    frame = Mock()
    frame.to_excel.side_effect = PermissionError("Synthetic locked file")
    log = Mock()
    with pytest.raises(PermissionError, match="Synthetic locked file"):
        export_excel(frame, tmp_path / "leads.xlsx", log=log)
    assert "Synthetic locked file" in [call.args[0] for call in log.call_args_list]
