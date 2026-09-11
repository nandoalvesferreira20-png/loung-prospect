from copy import deepcopy
from datetime import datetime, timedelta, timezone
from unittest.mock import Mock

from openpyxl import load_workbook

from core import scraper
from core.qualification.models import QualificationNote, QualificationResult
from core.qualification.serialization import qualification_result_to_columns


COLUMNS = ["Qualification Status", "Qualification Score", "Opportunity",
           "Qualification Reasons", "Qualification Evidence", "Qualification Limitations",
           "Rules Version", "Analyzed At"]


def test_serialization_is_readable_deterministic_and_does_not_mutate():
    result = QualificationResult(
        "qualified", score=80, opportunity="website",
        reasons=[QualificationNote("a", "Primeiro"), QualificationNote("b", "Segundo")],
        evidence=[QualificationNote("source", "Fonte consultada")],
        limitations=[QualificationNote("review", "Revisar")], rules_version="1.0.0",
        analyzed_at=datetime(2026, 1, 1, 12, 30, tzinfo=timezone(timedelta(hours=-3))),
    )
    original = deepcopy(result)
    actual = qualification_result_to_columns(result)
    assert list(actual) == COLUMNS
    assert list(actual.values()) == ["qualified", 80, "website", "a: Primeiro | b: Segundo",
                                     "source: Fonte consultada", "review: Revisar", "1.0.0",
                                     "2026-01-01T12:30:00-03:00"]
    assert actual == qualification_result_to_columns(result)
    assert result == original
    assert type(actual["Qualification Score"]) is int


def test_optional_values_and_empty_notes():
    result = qualification_result_to_columns(QualificationResult("error"))
    assert result["Qualification Status"] == "error"
    for column in ("Qualification Score", "Opportunity", "Rules Version", "Analyzed At"):
        assert result[column] is None
    for column in ("Qualification Reasons", "Qualification Evidence", "Qualification Limitations"):
        assert result[column] == ""


def test_excel_alignment_after_dedup_error_blanks_and_nonmutation(harness, lead_factory, monkeypatch, columns):
    records = [lead_factory(Empresa="A"), lead_factory(Empresa="A"), lead_factory(Empresa="B")]
    original = deepcopy(records)
    harness.links.return_value = ["a", "b", "c"]
    harness.extract.side_effect = records
    instant = datetime(2026, 1, 1, tzinfo=timezone.utc)
    results = [QualificationResult("qualified", score=80, opportunity="website", rules_version="1.0.0", analyzed_at=instant),
               QualificationResult("error", limitations=[QualificationNote("adapter_error", "Falha de adaptação")], analyzed_at=instant)]
    original_results = deepcopy(results)
    monkeypatch.setattr(scraper, "qualify_records", Mock(return_value=results))
    summary = harness.run(qualification_enabled=True)
    workbook = load_workbook(harness.output)
    try:
        rows = list(workbook.active.values)
        assert list(rows[0]) == columns + COLUMNS
        first, second = (dict(zip(rows[0], row)) for row in rows[1:])
        assert [first["Empresa"], second["Empresa"]] == ["A", "B"]
        assert first["Qualification Status"] == "qualified"
        assert type(first["Qualification Score"]) is int
        assert first["Qualification Score"] == 80
        assert first["Rules Version"] == "1.0.0"
        assert first["Analyzed At"] == instant.isoformat()
        assert second["Qualification Status"] == "error"
        assert second["Qualification Score"] is None
        assert second["Qualification Limitations"] == "adapter_error: Falha de adaptação"
    finally:
        workbook.close()
    assert records == original
    assert results == original_results
    assert [p["registro"] for p in summary["qualificacoes"]] == [original[0], original[2]]


def test_disabled_excel_content_matches_original(harness, lead_factory, columns):
    record = lead_factory()
    harness.extract.side_effect = [record]
    harness.run(qualification_enabled=False)
    workbook = load_workbook(harness.output)
    try:
        rows = list(workbook.active.values)
        assert list(rows[0]) == columns
        assert list(rows[1]) == [value if value != "" else None for value in record.values()]
    finally:
        workbook.close()
