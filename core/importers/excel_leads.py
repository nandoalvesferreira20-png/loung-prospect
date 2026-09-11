"""Import values from the first worksheet; formulas are never evaluated.

    Each successful insert commits independently through LeadRepository. Dedupe
    is an in-memory snapshot, not a concurrent uniqueness constraint.
"""
from dataclasses import dataclass, field
from pathlib import Path
import re

from openpyxl import load_workbook

MAPPING = {
    "Empresa": "empresa", "Cidade": "cidade", "Segmento": "segmento",
    "Telefone": "telefone", "WhatsApp": "whatsapp", "Email": "email",
    "Site": "site", "Endereço": "endereco", "Avaliação": "avaliacao",
    "Google Maps": "google_maps", "Observações": "observacoes",
    "Qualification Status": "qualification_status", "Qualification Score": "qualification_score",
    "Opportunity": "opportunity", "Qualification Reasons": "qualification_reasons",
    "Qualification Evidence": "qualification_evidence", "Qualification Limitations": "qualification_limitations",
    "Rules Version": "rules_version", "Analyzed At": "analyzed_at",
}


@dataclass(frozen=True)
class ImportIssue:
    row: int | None
    code: str
    message: str


@dataclass
class ImportSummary:
    total_rows: int = 0
    imported: int = 0
    skipped_duplicates: int = 0
    skipped_invalid: int = 0
    errors: list[ImportIssue] = field(default_factory=list)


def _text(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    return str(value)


def _normalized(value):
    return " ".join(str(value or "").lower().split())


def _key(data):
    # A Maps URL is opaque: exact comparison, not lowercase/path normalization.
    if data.get("google_maps"):
        return ("maps", data["google_maps"])
    name = _normalized(data.get("empresa"))
    address = _normalized(data.get("endereco"))
    if address:
        return ("address", name, address)
    phone = re.sub(r"[^0-9]", "", str(data.get("telefone") or ""))
    return ("phone", name, phone) if phone else None


def _score(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError("Qualification Score must be an integer from 0 to 100")
    if isinstance(value, str):
        if not re.fullmatch(r"\d+", value.strip()):
            raise ValueError("Qualification Score must be an integer from 0 to 100")
        value = int(value)
    if not isinstance(value, (int, float)) or not 0 <= value <= 100 or int(value) != value:
        raise ValueError("Qualification Score must be an integer from 0 to 100")
    return int(value)


def import_leads_from_excel(file_path, repository, *, default_responsible=None) -> ImportSummary:
    """Return file/row errors; do not touch the default DB implicitly.

    Header names are case-sensitive, trimmed. Empty sheets are no-ops. Nonempty
    sheets require Empresa. Entirely empty rows are ignored. Responsavel takes
    precedence over Responsável. An empty responsible cell falls back to default.
    Status Comercial is preserved when nonblank; otherwise status is Novo.
    Próxima Ação and Potencial (1-5) have no V1 database fields and are ignored.
    """
    summary = ImportSummary()
    try:
        if Path(file_path).suffix.lower() != ".xlsx":
            raise ValueError("Expected an .xlsx file")
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as error:
        summary.errors.append(ImportIssue(None, "file_error", str(error)))
        return summary
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = next(rows, ())
        if not any(value is not None for value in header):
            if any(any(value is not None for value in row) for row in rows):
                summary.errors.append(ImportIssue(1, "header_error", "Missing Empresa header"))
            return summary
        headers = [str(value).strip() if value is not None else "" for value in header]
        known = set(MAPPING) | {"Status Comercial", "Responsavel", "Responsável"}
        if "Empresa" not in headers or any(headers.count(name) > 1 for name in known):
            summary.errors.append(ImportIssue(1, "header_error", "Missing Empresa or duplicate known header"))
            return summary
        seen = {key for row in repository.list_leads() if (key := _key(dict(row))) is not None}
        for row_number, row in enumerate(rows, start=2):
            if all(value is None or (isinstance(value, str) and not value.strip()) for value in row):
                continue
            summary.total_rows += 1
            try:
                cells = dict(zip(headers, row))
                data = {target: _text(cells.get(source)) for source, target in MAPPING.items()}
                if not data["empresa"]:
                    raise ValueError("Empresa is required")
                data["qualification_score"] = _score(cells.get("Qualification Score"))
                data["status"] = _text(cells.get("Status Comercial")) or "Novo"
                responsible = cells.get("Responsavel") if "Responsavel" in cells else cells.get("Responsável")
                data["responsavel"] = _text(responsible) or _text(default_responsible)
                key = _key(data)
            except (ValueError, TypeError, OverflowError) as error:
                summary.skipped_invalid += 1
                summary.errors.append(ImportIssue(row_number, "invalid_row", str(error)))
                continue
            if key is not None and key in seen:
                summary.skipped_duplicates += 1
                continue
            try:
                repository.create_lead(data)
            except Exception as error:
                summary.skipped_invalid += 1
                summary.errors.append(ImportIssue(row_number, "write_error", str(error)))
                continue
            summary.imported += 1
            if key is not None:
                seen.add(key)
    except Exception as error:
        summary.errors.append(ImportIssue(None, "import_error", str(error)))
    finally:
        workbook.close()
    return summary
